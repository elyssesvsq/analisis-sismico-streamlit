import io
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter

# =====================================================================
# 1. FUNCIONES MATEMÁTICAS BASE (ENTREGABLE N° 1)
# =====================================================================
def calc_factor_C(T, Tp, TL):
    if T <= Tp:
        return 2.5
    elif Tp < T < TL:
        return 2.5 * (Tp / T)
    else:
        return 2.5 * ((Tp * TL) / (T ** 2))

def calc_pseudo_aceleracion(C, Z, U, S, R, g):
    return g * (Z * U * C * S) / R

def calc_frecuencia_estructural(T):
    return (2 * np.pi) / T

def calc_participacion_masa(X, m_vec, M_mat):
    return np.dot(X.T, m_vec) / np.dot(X.T, np.dot(M_mat, X))

def calc_fuerzas_y_cortantes(Sa, gamma, X, M_mat):
    F = np.dot(M_mat, Sa * gamma * X)
    V = np.cumsum(F)  # Acumulado desde el piso superior hacia la base
    return F, V

# =====================================================================
# 2. MOTOR INTEGRAL DE ANÁLISIS MODAL ESPECTRAL
# =====================================================================
def ejecutar_analisis_completo(pisos, masa, phi, T_modos, params):
    Z, U, S, R = params['Z'], params['U'], params['S'], params['R']
    Tp, TL, g = params['Tp'], params['TL'], params['g']
    
    M = np.diag(masa)
    n_modos = len(T_modos)
    
    C = np.array([calc_factor_C(t, Tp, TL) for t in T_modos])
    Sa = np.array([calc_pseudo_aceleracion(c, Z, U, S, R, g) for c in C])
    W = np.array([calc_frecuencia_estructural(t) for t in T_modos])
    Gamma = np.array([calc_participacion_masa(phi[:, i], masa, M) for i in range(n_modos)])
    
    F_modos = np.zeros_like(phi, dtype=float)
    V_modos = np.zeros_like(phi, dtype=float)
    for i in range(n_modos):
        F_modos[:, i], V_modos[:, i] = calc_fuerzas_y_cortantes(Sa[i], Gamma[i], phi[:, i], M)
        
    V_sum_abs = np.sum(np.abs(V_modos), axis=1)
    V_rcsc = np.sqrt(np.sum(V_modos ** 2, axis=1))
    Vrnc_h = 0.25 * V_sum_abs + 0.75 * V_rcsc
    Vreal = Vrnc_h * 0.75 * R
    
    df_modos = pd.DataFrame({
        'Modo': [f'Modo {i+1}' for i in range(n_modos)],
        'Periodo T (s)': T_modos,
        'Factor C': C,
        'Sa (m/s²)': Sa,
        'Frecuencia ω (rad/s)': W,
        'Participación Γ': Gamma,
        'Cortante Basal |V1| (tonf)': np.abs(V_modos[-1, :])
    })
    
    df_resultados = pd.DataFrame({
        'Piso': pisos,
        'Masa (t·s²/m)': masa,
        'F_M1 (tonf)': F_modos[:, 0], 'V_M1 (tonf)': V_modos[:, 0],
        'F_M2 (tonf)': F_modos[:, 1], 'V_M2 (tonf)': V_modos[:, 1],
        'F_M3 (tonf)': F_modos[:, 2], 'V_M3 (tonf)': V_modos[:, 2],
        'V_sum_abs (tonf)': V_sum_abs,
        'V_rcsc (tonf)': V_rcsc,
        'Vrnc_h (tonf)': Vrnc_h,
        'Vreal (tonf)': Vreal
    })
    
    return {
        'C': C, 'Sa': Sa, 'W': W, 'Gamma': Gamma,
        'F_modos': F_modos, 'V_modos': V_modos,
        'df_modos': df_modos, 'df_resultados': df_resultados
    }

# =====================================================================
# 3. FUNCIONES DE LOS 4 GRÁFICOS DE CÁLCULOS Y RESULTADOS
# =====================================================================
def grafico_1_espectro(T_modos, Sa_modos, params):
    """Gráfico 1: Espectro de Pseudo-aceleración de Diseño y Modos."""
    T_vec = np.linspace(0.01, 3.5, 300)
    C_vec = [calc_factor_C(t, params['Tp'], params['TL']) for t in T_vec]
    Sa_vec = [calc_pseudo_aceleracion(c, params['Z'], params['U'], params['S'], params['R'], params['g']) for c in C_vec]
    
    fig, ax = plt.subplots(figsize=(6, 4.5))
    ax.plot(T_vec, Sa_vec, color='#800000', lw=2.5, label='Espectro Sa(T)')
    colores = ['#1f77b4', '#ff7f0e', '#2ca02c']
    for i, (t, sa) in enumerate(zip(T_modos, Sa_modos)):
        ax.scatter(t, sa, color=colores[i], s=90, zorder=5, label=f'Modo {i+1} (T={t:.3f}s)')
        ax.vlines(t, 0, sa, color=colores[i], linestyle=':', alpha=0.8)
    ax.set_title('1. Espectro de Pseudo-Aceleración Inelástico', fontweight='bold')
    ax.set_xlabel('Periodo T (s)')
    ax.set_ylabel('Pseudo-aceleración Sa (m/s²)')
    ax.set_ylim(bottom=0)
    ax.grid(True, linestyle='--', alpha=0.5)
    ax.legend()
    plt.tight_layout()
    return fig

def grafico_2_fuerzas_modales(pisos, F_modos):
    """Gráfico 2: Distribución de Fuerzas Sísmicas F por piso."""
    pisos_con_base = np.append(pisos, 0)
    fig, ax = plt.subplots(figsize=(6, 4.5))
    for i in range(F_modos.shape[1]):
        F_plot = np.append(F_modos[:, i], 0)
        ax.plot(F_plot, pisos_con_base, marker='o', lw=2, label=f'Modo {i+1}')
    ax.axvline(0, color='black', lw=0.8)
    ax.set_title('2. Fuerzas Sísmicas F por Piso y Modo', fontweight='bold')
    ax.set_ylabel('Nivel / Piso')
    ax.set_xlabel('Fuerza Sísmica F (tonf)')
    ax.set_yticks(pisos_con_base)
    ax.grid(True, linestyle='--', alpha=0.5)
    ax.legend()
    plt.tight_layout()
    return fig

def grafico_3_cortantes_modales(pisos, V_modos):
    """Gráfico 3: Cortantes por Modo en cada nivel."""
    pisos_con_base = np.append(pisos, 0)
    fig, ax = plt.subplots(figsize=(6, 4.5))
    for i in range(V_modos.shape[1]):
        ax.barh(pisos - 0.5, V_modos[:, i], height=0.8, align='center', alpha=0.45, edgecolor='black', label=f'Modo {i+1}')
    ax.axvline(0, color='black', lw=0.8)
    ax.set_title('3. Fuerza Cortante V por Modo', fontweight='bold')
    ax.set_ylabel('Nivel / Piso')
    ax.set_xlabel('Fuerza Cortante V (tonf)')
    ax.set_yticks(pisos_con_base)
    ax.grid(True, linestyle='--', alpha=0.5)
    ax.legend()
    plt.tight_layout()
    return fig

def grafico_4_combinaciones(pisos, df_resultados):
    """Gráfico 4: Lienzo 2x2 de Combinaciones Normativas y Cortante Real."""
    pisos_con_base = np.append(pisos, 0)
    fig, axs = plt.subplots(2, 2, figsize=(10, 7.5))
    axs = axs.flatten()
    metrics = ['V_sum_abs (tonf)', 'V_rcsc (tonf)', 'Vrnc_h (tonf)', 'Vreal (tonf)']
    titles = ['Suma Absoluta (V_sum_abs)', 'Raíz C. Suma C. (V_rcsc)', 'Combinación Normativa (Vrnc_h)', 'Cortante Real de Diseño (Vreal)']
    colors = ['#1f77b4', '#ff7f0e', '#2ca02c', '#800000']

    for i in range(4):
        vals = df_resultados[metrics[i]].values
        bars = axs[i].barh(pisos - 0.5, vals, height=0.8, align='center', color=colors[i], alpha=0.7, edgecolor='black')
        axs[i].set_title(titles[i], fontweight='bold')
        axs[i].set_ylabel('Piso')
        axs[i].set_xlabel('Cortante (tonf)')
        axs[i].set_yticks(pisos_con_base)
        axs[i].grid(True, linestyle='--', alpha=0.5)
        for bar in bars:
            w = bar.get_width()
            axs[i].text(w * 0.98, bar.get_y() + bar.get_height()/2, f'{w:.1f}',
                        va='center', ha='right', color='white', fontweight='bold', fontsize=8)
    plt.tight_layout()
    return fig

# =====================================================================
# 4. EXPORTADOR EXCEL TÉCNICO MEJORADO
# =====================================================================
def generar_excel_en_memoria(params, df_modos, df_resultados):
    output = io.BytesIO()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Informe Sísmico"

    title_font = Font(bold=True, size=14, color="800000")
    sec_font = Font(bold=True, size=11, color="800000")
    header_font = Font(bold=True, color="FFFFFF", size=10)
    header_fill = PatternFill("solid", fgColor="800000")
    sub_fill = PatternFill("solid", fgColor="F2DCDB")
    border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    align_c = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws['A1'] = "INFORME TÉCNICO: ANÁLISIS DINÁMICO MODAL ESPECTRAL (NORMA E.030)"
    ws['A1'].font = title_font

    ws['A3'] = "1. PARÁMETROS NORMATIVOS DE ENTRADA"
    ws['A3'].font = sec_font
    param_keys = ['Z (Zona)', 'U (Uso)', 'S (Suelo)', 'R (Reducción)', 'Tp (s)', 'TL (s)', 'g (m/s²)']
    param_vals = [params['Z'], params['U'], params['S'], params['R'], params['Tp'], params['TL'], params['g']]
    for idx, (k, v) in enumerate(zip(param_keys, param_vals), start=1):
        c_head = ws.cell(row=4, column=idx, value=k)
        c_val = ws.cell(row=5, column=idx, value=v)
        c_head.font, c_head.fill, c_head.border, c_head.alignment = header_font, header_fill, border, align_c
        c_val.border, c_val.alignment, c_val.fill = border, align_c, sub_fill
        c_val.number_format = '0.00'

    ws['A7'] = "2. PROPIEDADES DINÁMICAS Y FACTORES POR MODO"
    ws['A7'].font = sec_font
    for col_idx, col_name in enumerate(df_modos.columns, start=1):
        cell = ws.cell(row=8, column=col_idx, value=col_name)
        cell.font, cell.fill, cell.border, cell.alignment = header_font, header_fill, border, align_c
    for r_idx, row in enumerate(df_modos.values, start=9):
        for c_idx, val in enumerate(row, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.border, cell.alignment = border, align_c
            if isinstance(val, (float, np.floating)):
                cell.number_format = '0.000'

    start_row = 14
    ws.cell(row=start_row - 1, column=1, value="3. RESULTADOS DE FUERZAS, CORTANTES MODALES Y COMBINACIONES").font = sec_font
    for col_idx, col_name in enumerate(df_resultados.columns, start=1):
        cell = ws.cell(row=start_row, column=col_idx, value=col_name)
        cell.font, cell.fill, cell.border, cell.alignment = header_font, header_fill, border, align_c

    n_pisos = len(df_resultados)
    for r_idx, row in enumerate(df_resultados.values, start=start_row + 1):
        for c_idx, val in enumerate(row, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=int(val) if c_idx == 1 else float(val))
            cell.border, cell.alignment = border, align_c
            if c_idx > 1:
                cell.number_format = '0.00'

    rango_vreal = f"L{start_row+1}:L{start_row+n_pisos}"
    regla_color = ColorScaleRule(start_type='min', start_color='FFF2CC',
                                 mid_type='percentile', mid_value=50, mid_color='F4B084',
                                 end_type='max', end_color='C00000')
    ws.conditional_formatting.add(rango_vreal, regla_color)

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 13)

    chart = BarChart()
    chart.type = "bar"
    chart.direction = "bar"
    chart.title = "Cortante Real de Diseño por Nivel (Vreal)"
    chart.x_axis.title = "Fuerza Cortante (tonf)"
    chart.y_axis.title = "Piso"
    data = Reference(ws, min_col=12, min_row=start_row, max_row=start_row + n_pisos)
    cats = Reference(ws, min_col=1, min_row=start_row + 1, max_row=start_row + n_pisos)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    ws.add_chart(chart, f"B{start_row + n_pisos + 2}")

    wb.save(output)
    output.seek(0)
    return output.getvalue()
